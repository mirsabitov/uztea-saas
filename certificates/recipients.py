"""Who gets a certificate: selected members, an uploaded Excel/CSV file, or pasted lines."""
import csv
import io
import re

from django.core.validators import validate_email
from django.core.exceptions import ValidationError

MAX_ROWS = 1000
ID_RE = re.compile(r'^\s*(?:UZ)?0*(\d{1,9})\s*$', re.IGNORECASE)
NAME_HEADERS = ('name', 'full name', 'full_name', 'ism', 'f.i.sh', 'fish', 'фио', 'имя', 'participant')
EMAIL_HEADERS = ('email', 'e-mail', 'mail', 'pochta', 'почта')
ID_HEADERS = ('membership id', 'membership_id', 'member id', 'id', "a'zolik id")


def _clean_email(value):
    value = (value or '').strip()
    if not value:
        return ''
    try:
        validate_email(value)
        return value.lower()
    except ValidationError:
        return None  # present but invalid


def _rows_from_table(table):
    """table: list of lists (first row may be a header)."""
    table = [[('' if c is None else str(c)).strip() for c in row] for row in table if any(c not in (None, '') for c in row)]
    if not table:
        return []
    header = [c.lower() for c in table[0]]
    idx = {'name': None, 'email': None, 'membership_id': None}
    for i, h in enumerate(header):
        if idx['name'] is None and h in NAME_HEADERS:
            idx['name'] = i
        elif idx['email'] is None and h in EMAIL_HEADERS:
            idx['email'] = i
        elif idx['membership_id'] is None and h in ID_HEADERS:
            idx['membership_id'] = i
    if idx['name'] is not None:
        body = table[1:]
    else:
        body = table  # no header: name, email, membership id
        idx = {'name': 0, 'email': 1, 'membership_id': 2}

    def cell(row, key):
        i = idx[key]
        return row[i] if i is not None and i < len(row) else ''

    return [{'name': cell(r, 'name'), 'email': cell(r, 'email'), 'membership_id': cell(r, 'membership_id')} for r in body]


def parse_upload(uploaded):
    name = (uploaded.name or '').lower()
    if name.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        wb = load_workbook(uploaded, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        table = [list(row) for row in ws.iter_rows(values_only=True)]
        return _rows_from_table(table[:MAX_ROWS + 1])
    data = uploaded.read()
    for encoding in ('utf-8-sig', 'cp1251', 'latin-1'):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=',;\t') if text.strip() else csv.excel
    return _rows_from_table(list(csv.reader(io.StringIO(text), dialect))[:MAX_ROWS + 1])


def parse_paste(text):
    rows = []
    for line in (text or '').splitlines():
        parts = [p.strip() for p in re.split(r'[\t;,]', line)]
        if any(parts):
            rows.append(parts)
    return _rows_from_table(rows[:MAX_ROWS])


def normalise(rows):
    """Validate rows and link them to memberships. Returns (good, problems)."""
    from app.models import Membership

    good, problems = [], []
    for n, row in enumerate(rows[:MAX_ROWS], start=1):
        name = (row.get('name') or '').strip()
        email = _clean_email(row.get('email'))
        membership = None
        match = ID_RE.match(str(row.get('membership_id') or ''))
        if match:
            membership = Membership.objects.filter(pk=int(match.group(1))).first()
        if not name and membership:
            name = membership.full_name
        if not name:
            problems.append(f'Row {n}: no name')
            continue
        if email is None:
            problems.append(f'Row {n} ({name}): invalid email "{row.get("email")}"')
            email = ''
        if membership is None and email:
            membership = Membership.objects.filter(email__iexact=email).order_by('-created_at').first()
        if not email and membership:
            email = (membership.email or '').lower()
        good.append({'name': name[:255], 'email': email, 'membership_id': membership.pk if membership else None})
    return good, problems


def from_members(ids):
    from app.models import Membership
    return [{'name': m.full_name, 'email': (m.email or '').lower(), 'membership_id': m.pk}
            for m in Membership.objects.filter(pk__in=ids).order_by('full_name')[:MAX_ROWS]]
